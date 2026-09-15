"""Livro como espelho — Slice 5 de docs/sdd/NARCISO_CANONICAL_SDD_v0.1.md (seção 18).

Extensões neutras do motor, todas opcionais:
- `check_visual_canon.py`: MIR-03 (`callback_changes`), `pair_with`,
  `display_assets` e a projeção por edição (`interior_plates`, `chapter_display`);
- `build_kdp_docx.py`: spread em verso+recto e recursos de exibição abaixo do título.

Render-QA (DOCX -> PDF) depende de LibreOffice no ambiente; aqui a paridade é
verificada pela estrutura: a metade esquerda abre seção em página par e a
direita na ímpar seguinte, o que obriga o Word a colocá-las frente a frente.

Execução (offline): .venv/Scripts/python.exe -m unittest tests.test_book_as_mirror -v
"""
from __future__ import annotations

import copy
import sys
import tempfile
import unittest
from io import BytesIO
from pathlib import Path

import yaml
from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from PIL import Image

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "engine" / "scripts"))

import build_kdp_docx as bd  # noqa: E402
import check_visual_canon as vc  # noqa: E402
from tests.test_illustration_system import _canon, _context  # noqa: E402
from tests.test_visual_narrative_rendering import _docx_content_hash, _write_minimal_runtime  # noqa: E402


def _dump(path: Path, data: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(yaml.safe_dump(data, allow_unicode=True, sort_keys=False), encoding="utf-8")


def _image_sections(doc) -> dict[str, int]:
    """título da imagem -> índice da seção. Um parágrafo com sectPr é o último
    da sua seção."""
    out, index = {}, 0
    for paragraph in doc.paragraphs:
        for doc_pr in paragraph._p.iter(qn("wp:docPr")):
            out[doc_pr.get("title")] = index
        ppr = paragraph._p.pPr
        if ppr is not None and ppr.find(qn("w:sectPr")) is not None:
            index += 1
    return out


def _blob(doc, shape) -> bytes:
    rid = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
    return doc.part.related_parts[rid].blob


class TestSpreadDocx(unittest.TestCase):

    def _runtime(self, tmp: Path, plate: dict) -> Path:
        runtime = _write_minimal_runtime(tmp, chapter_count=3)
        _dump(runtime / "layout" / "IMAGE_PLACEMENT.yaml", {"illustrations": [plate]})
        return runtime

    def test_spread_lands_on_verso_then_recto(self):
        plate = {"id": "IL-05", "chapter": 2, "placement": "AFTER", "anchor": "capítulo 2",
                 "pair_with": "SPREAD", "category": "HERO"}
        with tempfile.TemporaryDirectory() as tmp:
            runtime = self._runtime(Path(tmp), plate)
            Image.new("RGB", (200, 150), (10, 10, 10)).save(runtime / "images/approved/IL-05.jpg")
            self.assertEqual(bd.build(runtime), 0)
            doc = Document(runtime / "outputs" / "KDP_DRAFT.docx")
            sections = _image_sections(doc)
            verso, recto = sections["Ilustração IL-05 — verso"], sections["Ilustração IL-05 — recto"]
            self.assertEqual(recto, verso + 1)
            self.assertEqual(doc.sections[verso].start_type, WD_SECTION.EVEN_PAGE)
            self.assertEqual(doc.sections[recto].start_type, WD_SECTION.ODD_PAGE)
            self.assertEqual(doc.sections[recto + 1].start_type, WD_SECTION.NEW_PAGE)
            halves = [Image.open(BytesIO(_blob(doc, s))).size for s in doc.inline_shapes]
            self.assertEqual(halves, [(100, 150), (100, 150)])

    def test_explicit_halves_are_used_as_is(self):
        plate = {"id": "IL-05", "chapter": 2, "placement": "AFTER", "anchor": "capítulo 2", "pair_with": "SPREAD"}
        with tempfile.TemporaryDirectory() as tmp:
            runtime = self._runtime(Path(tmp), plate)
            files = []
            for side, shade in (("verso", 30), ("recto", 90)):
                path = runtime / f"images/approved/IL-05_{side}.jpg"
                Image.new("RGB", (90, 150), (shade, shade, shade)).save(path)
                files.append(path.read_bytes())
            self.assertEqual(bd.build(runtime), 0)
            doc = Document(runtime / "outputs" / "KDP_DRAFT.docx")
            self.assertEqual([_blob(doc, s) for s in doc.inline_shapes], files)

    def test_spread_before_opening_fails(self):
        plate = {"id": "IL-05", "chapter": 2, "placement": "OPEN", "pair_with": "SPREAD"}
        with tempfile.TemporaryDirectory() as tmp:
            runtime = self._runtime(Path(tmp), plate)
            with self.assertRaises(SystemExit) as raised:
                bd.build(runtime)
            self.assertIn("HINGE ou AFTER", str(raised.exception))

    def test_missing_spread_image_fails(self):
        plate = {"id": "IL-05", "chapter": 2, "placement": "AFTER", "anchor": "capítulo 2", "pair_with": "SPREAD"}
        with tempfile.TemporaryDirectory() as tmp:
            runtime = self._runtime(Path(tmp), plate)
            with self.assertRaises(SystemExit):
                bd.build(runtime)


class TestDisplayAssetsDocx(unittest.TestCase):

    def test_included_display_asset_goes_below_the_title(self):
        with tempfile.TemporaryDirectory() as tmp:
            runtime = _write_minimal_runtime(Path(tmp), chapter_count=3)
            (runtime / "images/display").mkdir(parents=True)
            Image.new("L", (400, 60), 255).save(runtime / "images/display/te_01.png")
            _dump(runtime / "layout/editions/kdp_paperback/EDITION_PLAN.yaml", {"chapter_display": [
                {"chapter": 1, "display_asset": "TE", "kind": "TITLE_ECHO", "asset": "images/display/te_01.png",
                 "status": "INCLUDED"},
                {"chapter": 2, "display_asset": "TE", "kind": "TITLE_ECHO", "asset": "images/display/te_02.png",
                 "status": "INCLUDED"},
                {"chapter": 3, "display_asset": "TE", "kind": "TITLE_ECHO", "asset": None, "status": "OMITTED"},
            ]})
            self.assertEqual(bd.build(runtime), 0)
            doc = Document(runtime / "outputs" / "KDP_DRAFT.docx")
            titles = [s._inline.docPr.get("title") for s in doc.inline_shapes]
            self.assertEqual([t for t in titles if t.startswith("Exibição")], ["Exibição TE — capítulo 01"])
            paragraphs = doc.paragraphs
            index = next(i for i, p in enumerate(paragraphs)
                         if any(d.get("title", "").startswith("Exibição") for d in p._p.iter(qn("wp:docPr"))))
            self.assertEqual(paragraphs[index - 1].style.name, "Heading 1")
            self.assertEqual(paragraphs[index + 1].style.name, "Body First")

    def test_plan_without_display_keeps_docx_identical(self):
        hashes = []
        for with_plan in (False, True):
            with tempfile.TemporaryDirectory() as tmp:
                runtime = _write_minimal_runtime(Path(tmp), chapter_count=2)
                if with_plan:
                    _dump(runtime / "layout/editions/kdp_paperback/EDITION_PLAN.yaml", {"finish_intents": []})
                self.assertEqual(bd.build(runtime), 0)
                hashes.append(_docx_content_hash(runtime / "outputs" / "KDP_DRAFT.docx"))
        self.assertEqual(hashes[0], hashes[1])


def _mirror_canon() -> dict:
    canon = _canon()
    plate = next(c for c in canon["compositions"] if c["id"] == "IL-03")
    plate["illustration"]["pair_with"] = "SPREAD"
    canon["display_assets"] = [
        {"id": "TE", "kind": "TITLE_ECHO", "chapters": [3, 4], "asset": "images/display/te_{chapter:02d}.png"},
        {"id": "ORN", "kind": "ORNAMENT", "chapters": [1], "asset": "images/display/orn.png", "omit_in": ["kdp_hardcover"]},
    ]
    return canon


class TestMirrorRules(unittest.TestCase):

    def findings(self, mutate=None) -> list[dict]:
        canon = _mirror_canon()
        if mutate:
            mutate(canon)
        context = _context()
        return vc.check_illustrations(canon, context) + vc.check_display_assets(canon, context)

    def assertCategory(self, mutate, category, severity):
        found = [f for f in self.findings(mutate) if f["category"] == category]
        self.assertTrue(found, f"{category} não emitido: {[f['category'] for f in self.findings(mutate)]}")
        self.assertIn(severity, {f["severity"] for f in found})

    @staticmethod
    def block(canon, cid):
        return next(c for c in canon["compositions"] if c["id"] == cid)["illustration"]

    def test_valid_mirror_canon(self):
        self.assertEqual(self.findings(), [])

    def test_callback_without_change(self):
        self.assertCategory(lambda c: self.block(c, "IL-03").pop("callback_changes"), "CALLBACK_WITHOUT_CHANGE", "MEDIUM")

    def test_callback_overloaded(self):
        self.assertCategory(lambda c: self.block(c, "IL-03").update(callback_changes=[
            {"what": "a", "impossible": True}, {"what": "b", "impossible": True}]), "CALLBACK_OVERLOADED", "MEDIUM")

    def test_spread_with_open_placement(self):
        self.assertCategory(lambda c: self.block(c, "IL-03").update(placement="OPEN"), "SPREAD_PLACEMENT_INVALID", "HIGH")

    def test_pair_with_other_chapter(self):
        self.assertCategory(lambda c: self.block(c, "IL-02").update(pair_with="IL-01"), "PAIR_WITH_UNRESOLVED", "HIGH")

    def test_display_asset_invalid_kind(self):
        self.assertCategory(lambda c: c["display_assets"][0].update(kind="MIRRORED_BODY"), "DISPLAY_ASSET_INVALID", "HIGH")

    def test_display_asset_chapter_out_of_range(self):
        self.assertCategory(lambda c: c["display_assets"][0].update(chapters=[99]), "DISPLAY_ASSET_INVALID", "HIGH")

    def test_display_asset_overlap(self):
        self.assertCategory(lambda c: c["display_assets"].append(
            {"id": "TE2", "kind": "TITLE_ECHO", "chapters": [4], "asset": "x.png"}), "DISPLAY_ASSET_INVALID", "HIGH")

    def test_display_asset_bad_placeholder(self):
        self.assertCategory(lambda c: c["display_assets"][0].update(asset="te_{id}.png"), "DISPLAY_ASSET_INVALID", "HIGH")

    def test_display_asset_unknown_target(self):
        self.assertCategory(lambda c: c["display_assets"][1].update(omit_in=["web"]), "DISPLAY_ASSET_INVALID", "HIGH")


class TestMirrorProjection(unittest.TestCase):

    def test_kindle_sequences_spread_and_omits_display(self):
        plan = vc.project_mirror_manifestation(_mirror_canon(), "kindle_ebook")
        layouts = {p["composition"]: p["layout"] for p in plan["interior_plates"]}
        self.assertEqual(layouts, {"IL-01": "SINGLE", "IL-02": "SINGLE", "IL-03": "SEQUENTIAL"})
        self.assertEqual({e["status"] for e in plan["chapter_display"]}, {"OMITTED"})
        self.assertTrue(all(e["asset"] is None for e in plan["chapter_display"]))

    def test_print_keeps_spread_and_formats_assets(self):
        plan = vc.project_mirror_manifestation(_mirror_canon(), "kdp_paperback")
        self.assertEqual(next(p for p in plan["interior_plates"] if p["composition"] == "IL-03")["layout"], "SPREAD")
        self.assertEqual([(e["chapter"], e["asset"]) for e in plan["chapter_display"]],
                         [(1, "images/display/orn.png"), (3, "images/display/te_03.png"), (4, "images/display/te_04.png")])

    def test_omit_in_is_honored(self):
        plan = vc.project_mirror_manifestation(_mirror_canon(), "kdp_hardcover")
        ornament = next(e for e in plan["chapter_display"] if e["display_asset"] == "ORN")
        self.assertEqual((ornament["status"], ornament["asset"]), ("OMITTED", None))

    def test_canon_without_plates_or_display_projects_nothing(self):
        canon = {"compositions": [{"id": "COMP-X", "surface": "FRONT_COVER"}]}
        for target in vc.EDITION_TARGETS:
            self.assertEqual(vc.project_mirror_manifestation(canon, target), {})

    def test_projection_is_deterministic(self):
        canon = _mirror_canon()
        self.assertEqual(vc.project_mirror_manifestation(canon, "collector"),
                         vc.project_mirror_manifestation(copy.deepcopy(canon), "collector"))


if __name__ == "__main__":
    unittest.main()
