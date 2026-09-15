"""Testes dos hooks de renderização existentes — Slice 6 da capability
`BEA_HALDEN_VISUAL_NARRATIVE_SYSTEM`
(ver docs/sdd/BEA_HALDEN_VISUAL_NARRATIVE_SYSTEM_SDD_v0.1.md, seção 31).

Cobre `engine/scripts/build_cover_and_stories.py` (`typography_bindings`
opcional em `load_font`, seção 17.4/D-V2) e `engine/scripts/build_kdp_docx.py`
(sigil de abertura de capítulo a partir de `layout/editions/<target>/
EDITION_PLAN.yaml`, seção 18.4) — os dois únicos scripts de renderização que
esta capability estende, sem mudar o grafo do compositor ("Hooks de
renderização existentes (sem mudança padrão)").

100% offline, imagens de teste geradas em diretório temporário com Pillow
(D10/29 da SDD): `.venv/Scripts/python.exe -m unittest
tests.test_visual_narrative_rendering -v`.
"""
from __future__ import annotations

import hashlib
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

import yaml
from docx import Document
from PIL import Image

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "engine" / "scripts"))

import build_cover_and_stories as bc  # noqa: E402
import build_kdp_docx as bd  # noqa: E402


def _write_minimal_runtime(tmp: Path, chapter_count: int = 2) -> Path:
    """Runtime sintético mínimo — só o que `build_cover_and_stories.py` e
    `build_kdp_docx.py` realmente leem. Nenhum canon visual envolvido: o
    objetivo destes testes é a extensão do RENDERIZADOR, não o validador
    (já coberto em `tests/test_visual_canon.py`)."""
    runtime = tmp / "rt"
    (runtime / "book").mkdir(parents=True)
    book_spec = {
        "apiVersion": "pedroarte.livingbooks/v1", "kind": "BookSpec",
        "metadata": {"slug": "rendering_test", "title": "Livro de Teste",
                     "author": "Autora de Teste", "language": "pt-BR",
                     "chapter_count": chapter_count},
        "spec": {"chapter_titles": {i: f"Capítulo {i}" for i in range(1, chapter_count + 1)}},
    }
    (runtime / "book" / "BOOK_SPEC.yaml").write_text(
        yaml.safe_dump(book_spec, allow_unicode=True), encoding="utf-8")

    manuscript = "\n\n---\n\n".join(
        f"# {i}. Capítulo {i}\n\nUm parágrafo de prosa neutra para o capítulo {i}."
        for i in range(1, chapter_count + 1)
    )
    (runtime / "manuscript" / "final").mkdir(parents=True)
    (runtime / "manuscript" / "final" / "MANUSCRIPT_FINAL_PTBR.md").write_text(
        manuscript, encoding="utf-8")

    (runtime / "images" / "approved").mkdir(parents=True)
    for i in range(1, chapter_count + 1):
        Image.new("RGB", (100, 150), (20, 20, 20)).save(
            runtime / "images" / "approved" / f"chapter_{i:02d}.jpg")
    return runtime


def _docx_content_hash(path: Path) -> str:
    """Hash do CONTEÚDO de cada parte do .docx (zip de XML), ignorando os
    timestamps que o zip embute por entrada — dois saves de um mesmo
    documento nunca são byte-idênticos por causa deles, mesmo com conteúdo
    logicamente idêntico (verificado empiricamente antes de escrever este
    teste, mesma disciplina de calibração do Slice 4)."""
    h = hashlib.sha256()
    with zipfile.ZipFile(path) as z:
        for name in sorted(z.namelist()):
            h.update(name.encode("utf-8"))
            h.update(z.read(name))
    return h.hexdigest()


class TestFontBindingExtension(unittest.TestCase):
    """D-V2 / seção 17.4 da SDD: `load_font()` aceita `typography_bindings`
    opcionais sem mudar o comportamento padrão."""

    @classmethod
    def setUpClass(cls):
        source, _ = bc.resolve_font("sans")
        if source is None:
            raise unittest.SkipTest("nenhuma fonte TrueType disponível neste ambiente")
        cls.stand_in_font = source

    def test_role_binding_used_when_embeddable_and_present(self):
        with tempfile.TemporaryDirectory() as tmp:
            font_path = Path(tmp) / "custom.ttf"
            font_path.write_bytes(self.stand_in_font.read_bytes())
            bindings = {"TITLE": {"family": "Custom", "file": str(font_path),
                                  "license": "OFL", "embeddable": True}}
            font = bc.load_font("serif_bold", 40, role="TITLE", typography_bindings=bindings)
            self.assertEqual(Path(font.path).resolve(), font_path.resolve())

    def test_non_embeddable_binding_falls_back_to_default(self):
        bindings = {"TITLE": {"file": str(self.stand_in_font), "embeddable": False}}
        font = bc.load_font("serif_bold", 40, role="TITLE", typography_bindings=bindings)
        default = bc.load_font("serif_bold", 40)
        self.assertEqual(getattr(font, "path", None), getattr(default, "path", None))

    def test_missing_binding_file_falls_back_to_default(self):
        bindings = {"TITLE": {"file": "does/not/exist.ttf", "embeddable": True}}
        font = bc.load_font("serif_bold", 40, role="TITLE", typography_bindings=bindings)
        default = bc.load_font("serif_bold", 40)
        self.assertEqual(getattr(font, "path", None), getattr(default, "path", None))

    def test_binding_for_a_different_role_is_not_applied(self):
        bindings = {"AUTHOR": {"file": str(self.stand_in_font), "embeddable": True}}
        font = bc.load_font("serif_bold", 40, role="TITLE", typography_bindings=bindings)
        default = bc.load_font("serif_bold", 40)
        self.assertEqual(getattr(font, "path", None), getattr(default, "path", None))

    def test_no_role_is_a_pure_no_op_even_with_bindings_present(self):
        bindings = {"TITLE": {"file": str(self.stand_in_font), "embeddable": True}}
        font = bc.load_font("serif_bold", 40, typography_bindings=bindings)
        default = bc.load_font("serif_bold", 40)
        self.assertEqual(getattr(font, "path", None), getattr(default, "path", None))


class TestChapterOpenerDocxInsertion(unittest.TestCase):
    """Seção 18.4 da SDD: sigil de abertura inserido acima de `CAPÍTULO N`
    quando `layout/editions/<target>/EDITION_PLAN.yaml` declara
    `chapter_openers[]` com um asset que existe no runtime."""

    def test_no_plan_means_zero_sigils_and_unaffected_build(self):
        with tempfile.TemporaryDirectory() as tmp:
            runtime = _write_minimal_runtime(Path(tmp), chapter_count=2)
            self.assertEqual(bd.build(runtime), 0)
            doc = Document(runtime / "outputs" / "KDP_DRAFT.docx")
            self.assertEqual(len(doc.inline_shapes), 2)  # só as 2 imagens de capítulo

    def test_plan_with_asset_inserts_sigil_only_where_the_file_exists(self):
        with tempfile.TemporaryDirectory() as tmp:
            runtime = _write_minimal_runtime(Path(tmp), chapter_count=3)
            sigil_dir = runtime / "images" / "sigils" / "SIG-TEST"
            sigil_dir.mkdir(parents=True)
            Image.new("1", (200, 200), 1).save(sigil_dir / "STATE1.png")

            plan_dir = runtime / "layout" / "editions" / "kdp_paperback"
            plan_dir.mkdir(parents=True)
            plan = {
                "target": "kdp_paperback",
                "chapter_openers": [
                    {"chapter": 1, "composition": "COMP-OPENER", "element": "SIG-TEST",
                     "state": "STATE1", "render_mode": "RASTER_1BIT",
                     "asset": "images/sigils/SIG-TEST/STATE1.png"},
                    # capítulo 2: asset declarado mas ausente do runtime —
                    # ignorado, não falha o build (18.4: reforço opcional).
                    {"chapter": 2, "composition": "COMP-OPENER", "element": "SIG-TEST",
                     "state": "STATE1", "render_mode": "RASTER_1BIT",
                     "asset": "images/sigils/SIG-TEST/MISSING.png"},
                    # capítulo 3: sem entrada no plano.
                ],
            }
            (plan_dir / "EDITION_PLAN.yaml").write_text(
                yaml.safe_dump(plan, allow_unicode=True), encoding="utf-8")

            self.assertEqual(bd.build(runtime), 0)
            doc = Document(runtime / "outputs" / "KDP_DRAFT.docx")
            self.assertEqual(len(doc.inline_shapes), 4)  # 3 imagens + 1 sigil
            titles = [s._inline.docPr.get("title") for s in doc.inline_shapes]
            self.assertIn("Sigil de abertura — capítulo 01", titles)
            self.assertNotIn("Sigil de abertura — capítulo 02", titles)
            self.assertNotIn("Sigil de abertura — capítulo 03", titles)

    def test_hardcover_plan_is_read_when_paperback_plan_is_absent(self):
        with tempfile.TemporaryDirectory() as tmp:
            runtime = _write_minimal_runtime(Path(tmp), chapter_count=1)
            sigil_dir = runtime / "images" / "sigils" / "SIG-TEST"
            sigil_dir.mkdir(parents=True)
            Image.new("1", (200, 200), 1).save(sigil_dir / "STATE1.png")

            plan_dir = runtime / "layout" / "editions" / "kdp_hardcover"
            plan_dir.mkdir(parents=True)
            plan = {"chapter_openers": [
                {"chapter": 1, "element": "SIG-TEST", "state": "STATE1",
                 "asset": "images/sigils/SIG-TEST/STATE1.png"},
            ]}
            (plan_dir / "EDITION_PLAN.yaml").write_text(
                yaml.safe_dump(plan, allow_unicode=True), encoding="utf-8")

            self.assertEqual(bd.build(runtime), 0)
            doc = Document(runtime / "outputs" / "KDP_DRAFT.docx")
            self.assertEqual(len(doc.inline_shapes), 2)  # 1 imagem + 1 sigil


class TestINVVN02Regression(unittest.TestCase):
    """INV-VN-02 (seção 28.3/31 da SDD): sem os arquivos derivados desta
    capability, `build_cover_and_stories.py` e `build_kdp_docx.py` produzem
    saída idêntica — determinística entre execuções independentes."""

    def test_cover_hash_is_identical_across_independent_runs(self):
        with tempfile.TemporaryDirectory() as tmp1, tempfile.TemporaryDirectory() as tmp2:
            digests = []
            for tmp in (tmp1, tmp2):
                runtime = _write_minimal_runtime(Path(tmp))
                book = bc.load_book_metadata(runtime)
                design = bc.merged_design(runtime)  # sem media/MEDIA_DESIGN.yaml
                self.assertNotIn("typography_bindings", design)
                _, digest = bc.build_cover(runtime, book, design, use_base=False)
                digests.append(digest)
            self.assertEqual(digests[0], digests[1])

    def test_docx_content_is_identical_across_independent_runs_without_plan(self):
        hashes = []
        for _ in range(2):
            with tempfile.TemporaryDirectory() as tmp:
                runtime = _write_minimal_runtime(Path(tmp), chapter_count=2)
                self.assertEqual(bd.build(runtime), 0)
                hashes.append(_docx_content_hash(runtime / "outputs" / "KDP_DRAFT.docx"))
        self.assertEqual(hashes[0], hashes[1])

    def test_typography_bindings_absent_takes_the_exact_old_code_path(self):
        # `design.get("typography_bindings")` é `None` sem projeção — o mesmo
        # valor que todo caller passava implicitamente antes do Slice 6.
        with tempfile.TemporaryDirectory() as tmp:
            runtime = _write_minimal_runtime(Path(tmp))
            design = bc.merged_design(runtime)
            self.assertIsNone(design.get("typography_bindings"))
            with_role = bc.load_font("serif_bold", 40, role="TITLE",
                                     typography_bindings=design.get("typography_bindings"))
            without_role = bc.load_font("serif_bold", 40)
            self.assertEqual(getattr(with_role, "path", None), getattr(without_role, "path", None))


if __name__ == "__main__":
    unittest.main()
